"""The curation worksheet: ranking, suggestions, and the round trip."""
import csv
import shutil

import pytest
from openpyxl import load_workbook

from quickquote.reference.curate import (
    INGREDIENT_SHEET,
    LOOKUP_SHEET,
    PACKAGING_SHEET,
    PRICING_SHEET,
    _changed,
    apply,
    build_rows,
    export,
    suggest_potency,
)
from quickquote.reference.loader import (
    PoRow,
    is_component_part,
    load_reference_data,
)


@pytest.fixture
def data_dir(tmp_path, reference):
    """A writable copy of the reference tables."""
    target = tmp_path / "reference"
    target.mkdir()
    for path in reference.source_dir.glob("*.csv"):
        shutil.copy(path, target / path.name)
    return target


@pytest.fixture
def worksheet(tmp_path, data_dir):
    path = tmp_path / "curation.xlsx"
    export(load_reference_data(data_dir), path)
    return path


def _po(part, description, uom="KG", spend=0.0):
    from datetime import date

    return PoRow(part_number=part, description=description, uom=uom,
                 latest_unit_cost=1.0, min_unit_cost_ever=1.0, max_unit_cost_ever=1.0,
                 latest_po_date=date(2026, 8, 1), latest_vendor="",
                 unique_vendor_count=0, po_count=1, total_spend=spend)


class TestSuggestedPotency:
    def test_a_stated_standardisation_is_offered(self):
        value, basis = suggest_potency("TURMERIC EG EXT. 95% ALL NAT")
        assert value == pytest.approx(0.95)
        assert "95%" in basis

    def test_the_basis_names_the_text_it_came_from(self):
        _, basis = suggest_potency("MILK THISTLE EXTRACT 80% UV")
        assert "description reads" in basis and "80%" in basis

    def test_a_description_with_no_percentage_suggests_nothing(self):
        assert suggest_potency("KSM-66 ASHWAGANDHA EXTRACT") == (None, "")

    def test_an_implausible_percentage_is_ignored(self):
        assert suggest_potency("BLEND 250% OVERAGE")[0] is None

    def test_the_suggestion_never_reaches_the_potency_column(self, reference):
        """A suggestion a person accepts is data; one applied silently is invention.

        Red yeast rice is the case that matters: its potency depends on
        monacolin standardisation, so the tables deliberately leave it unknown.
        A description stating "5%" must not quietly fill that in.
        """
        import copy

        catalogue = copy.deepcopy(reference)
        catalogue.identity_index = None
        catalogue.po_rows = [
            _po("RM-RYR", "RED YEAST RICE EXTRACT 5% MONACOLIN", spend=100.0)
        ]
        row = build_rows(catalogue)[0][0]

        assert row.identity == "red yeast rice"
        assert row.suggested_potency == pytest.approx(0.05)
        assert "5%" in row.suggestion_basis
        assert row.potency is None          # the column a person fills stays empty

    def test_a_curated_potency_is_shown_as_context(self, reference):
        import copy

        catalogue = copy.deepcopy(reference)
        catalogue.identity_index = None
        catalogue.po_rows = [
            _po("RM-TUR", "Turmeric Extract 95% Curcuminoids", spend=100.0)
        ]
        row = build_rows(catalogue)[0][0]
        entry = reference.identity_entry("turmeric extract")
        assert row.potency == entry.default_potency


class TestRanking:
    def test_rows_are_ordered_by_purchasing_spend(self, reference):
        import copy

        catalogue = copy.deepcopy(reference)
        catalogue.identity_index = None
        catalogue.po_rows = [
            _po("RM-A", "Coenzyme Q10 Ubiquinone USP Powder", spend=10.0),
            _po("RM-B", "Ashwagandha Root Powder Organic", spend=500.0),
            _po("RM-C", "Rice Flour Organic Fine Mesh", spend=100.0),
        ]
        ingredients, _ = build_rows(catalogue)
        assert [row.row.part_number for row in ingredients] == ["RM-B", "RM-C", "RM-A"]

    def test_cumulative_share_reaches_one(self, reference):
        import copy

        catalogue = copy.deepcopy(reference)
        catalogue.identity_index = None
        catalogue.po_rows = [
            _po("RM-A", "Coenzyme Q10 Ubiquinone USP Powder", spend=750.0),
            _po("RM-B", "Ashwagandha Root Powder Organic", spend=250.0),
        ]
        ingredients, _ = build_rows(catalogue)
        assert ingredients[0].cumulative == pytest.approx(0.75)
        assert ingredients[-1].cumulative == pytest.approx(1.0)


class TestComponentClassification:
    def test_per_thousand_pricing_means_a_component(self):
        """Capsule shells are components, not raw materials."""
        assert is_component_part(_po("RECVS00", "SIZE 00 VEGETABLE CAPSULE", uom="M"))

    def test_each_pricing_means_a_component(self):
        assert is_component_part(_po("KTTP225", "225cc WHITE HDPE BOTTLE", uom="EA"))

    def test_weight_pricing_means_a_raw_material(self):
        assert not is_component_part(_po("RAGLU1", "L-GLUTAMINE", uom="KG"))

    def test_shells_land_in_the_packaging_queue(self, reference):
        import copy

        catalogue = copy.deepcopy(reference)
        catalogue.identity_index = None
        catalogue.po_rows = [
            _po("RECVS00", "SIZE 00 VEGETABLE CAPSULE", uom="M", spend=100.0),
            _po("RAGLU1", "L-Glutamine", uom="KG", spend=100.0),
        ]
        ingredients, packaging = build_rows(catalogue)
        assert [row.row.part_number for row in packaging] == ["RECVS00"]
        assert [row.row.part_number for row in ingredients] == ["RAGLU1"]


class TestWorksheet:
    def test_it_carries_a_tab_for_each_owner(self, worksheet):
        book = load_workbook(worksheet)
        visible = [name for name in book.sheetnames
                   if book[name].sheet_state == "visible"]
        assert visible == ["Ingredients", "Packaging", "Machines",
                           "Process Steps", "Margins"]

    def test_the_process_steps_tab_names_every_gap(self, worksheet, reference):
        """Operations needs one place showing what stops a line being quotable."""
        sheet = load_workbook(worksheet)["Process Steps"]
        header = [cell.value for cell in sheet[1]]
        rows = [dict(zip(header, row))
                for row in sheet.iter_rows(min_row=2, values_only=True)
                if row[0] and row[1]]      # the trailing note row has no step
        forms = {row["Dosage Form"] for row in rows}
        assert forms >= {"capsule", "tablet", "gummy"}

        gummy = [row for row in rows if row["Dosage Form"] == "gummy"]
        assert any(row["Status"] == "No work centre assigned" for row in gummy)
        assert all(row["Core Step"] in ("core", "ancillary") for row in rows)

    def test_the_overage_column_offers_every_class(self, worksheet, reference):
        """Excel truncates an inline list at 255 characters, which is fewer
        classes than R&D's guideline carries, so the list is a range."""
        book = load_workbook(worksheet)
        rule = book[INGREDIENT_SHEET].data_validations.dataValidation[0]
        assert rule.formula1.startswith(f"'{LOOKUP_SHEET}'!")

        offered = {row[0] for row in book[LOOKUP_SHEET].iter_rows(
            min_row=2, max_col=1, values_only=True) if row[0]}
        assert offered == set(reference.overage)
        assert book[LOOKUP_SHEET].sheet_state == "hidden"


class TestRoundTrip:
    def test_an_untouched_worksheet_changes_nothing(self, worksheet, data_dir):
        """The export pre-fills current values as context; re-applying them is
        not a confirmation and must not stamp defaults as reviewed."""
        assert not any(apply(worksheet, data_dir).values())

    def test_a_filled_potency_reaches_the_engine(self, worksheet, data_dir):
        book = load_workbook(worksheet)
        sheet = book[INGREDIENT_SHEET]
        header = [cell.value for cell in sheet[1]]
        identity = sheet.cell(row=2, column=header.index("Identity") + 1).value
        sheet.cell(row=2, column=header.index("Potency *") + 1).value = 0.42
        book.save(worksheet)

        changed = apply(worksheet, data_dir)
        assert changed["potency"] == 1
        assert load_reference_data(data_dir).identity_entry(identity).default_potency == 0.42

    def test_a_filled_density_is_recorded_as_confirmed(self, worksheet, data_dir):
        book = load_workbook(worksheet)
        sheet = book[INGREDIENT_SHEET]
        header = [cell.value for cell in sheet[1]]
        identity = sheet.cell(row=2, column=header.index("Identity") + 1).value
        sheet.cell(row=2, column=header.index("Bulk Density g/mL *") + 1).value = 0.62
        book.save(worksheet)

        apply(worksheet, data_dir)
        rows = {row["key"]: row for row in csv.DictReader((data_dir / "bulk_density.csv").open())}
        assert float(rows[identity]["bulk_density_g_ml"]) == 0.62
        assert "Confirmed" in rows[identity]["notes"]

    def test_a_material_overage_override_reaches_the_engine(self, worksheet, data_dir):
        """One unusual material can be set without moving its whole class."""
        book = load_workbook(worksheet)
        sheet = book[INGREDIENT_SHEET]
        header = [cell.value for cell in sheet[1]]
        identity = sheet.cell(row=2, column=header.index("Identity") + 1).value
        sheet.cell(row=2, column=header.index("Overage % override") + 1).value = 18
        book.save(worksheet)

        assert apply(worksheet, data_dir)["overage_override"] == 1
        entry = load_reference_data(data_dir).identity_entry(identity)
        assert entry.overage_pct == 18

    def test_a_run_rate_reaches_the_engine(self, worksheet, data_dir):
        book = load_workbook(worksheet)
        sheet = book["Process Steps"]
        header = [cell.value for cell in sheet[1]]
        row = next(index for index in range(2, sheet.max_row + 1)
                   if sheet.cell(row=index, column=header.index("Work Centre") + 1).value
                   == "Tablet press")
        sheet.cell(row=row, column=header.index("Units per Hour *") + 1).value = 42000
        book.save(worksheet)

        assert apply(worksheet, data_dir)["run_rates"] == 1
        rate = load_reference_data(data_dir).run_rates["Tablet press"]
        assert rate.units_per_hour == 42000 and rate.confirmed

    def test_a_margin_target_reaches_the_engine(self, worksheet, data_dir):
        book = load_workbook(worksheet)
        book[PRICING_SHEET].cell(row=2, column=3).value = 32.5
        book.save(worksheet)

        assert apply(worksheet, data_dir)["margins"] == 1
        pricing = load_reference_data(data_dir).pricing
        contract = next(item for item in pricing if item.channel == "contract")
        assert contract.target_margin_pct == 32.5
        assert "Confirmed by Finance" in contract.notes

    def test_a_blank_cell_never_clears_a_curated_value(self, worksheet, data_dir):
        before = (data_dir / "ingredient_identity.csv").read_text()
        book = load_workbook(worksheet)
        sheet = book[INGREDIENT_SHEET]
        header = [cell.value for cell in sheet[1]]
        for row in range(2, min(sheet.max_row, 40) + 1):
            sheet.cell(row=row, column=header.index("Potency *") + 1).value = None
            sheet.cell(row=row, column=header.index("Overage Class *") + 1).value = None
        book.save(worksheet)

        apply(worksheet, data_dir)
        assert (data_dir / "ingredient_identity.csv").read_text() == before

    def test_a_packaging_role_reaches_the_engine(self, worksheet, data_dir):
        book = load_workbook(worksheet)
        sheet = book[PACKAGING_SHEET]
        header = [cell.value for cell in sheet[1]]
        identity = sheet.cell(row=2, column=header.index("Identity") + 1).value
        if not identity:
            pytest.skip("no resolved packaging identity in the demo set")
        sheet.cell(row=2, column=header.index("Bottles per Purchased Unit *") + 1).value = 24
        book.save(worksheet)

        apply(worksheet, data_dir)
        entry = load_reference_data(data_dir).identity_entry(identity, packaging=True)
        assert entry.units_per_container == 24


class TestChangeDetection:
    @pytest.mark.parametrize("new,current,expected", [
        (None, "1.0", False),
        ("", "1.0", False),
        (0.95, "0.95", False),
        (0.95, "1.0", True),
        ("botanical_extract", "botanical_extract", False),
        ("botanical_extract", "default", True),
        (0.5, None, True),
    ])
    def test_only_a_real_answer_counts(self, new, current, expected):
        assert _changed(new, current) is expected
