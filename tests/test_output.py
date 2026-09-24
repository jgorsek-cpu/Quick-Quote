"""Workbook and PDF structure."""
import io

import pytest
from openpyxl import load_workbook

from quickquote.config import EXCLUDED_DISPLAY
from quickquote.engine.pipeline import run_pipeline
from quickquote.output.pdf import build_quote_pdf
from quickquote.output.workbook import build_workbook

from .conftest import AS_OF

EXPECTED_TABS = [
    "Product Summary", "BOM - Raw Materials", "Packaging",
    "Cost Summary", "Review Flags", "Customer Quote Summary",
]


@pytest.fixture
def result(heart_health, reference):
    return run_pipeline(heart_health, reference, as_of=AS_OF)


@pytest.fixture
def book(result):
    return load_workbook(io.BytesIO(build_workbook(result)))


def test_every_quote_produces_the_same_six_tabs(book):
    assert book.sheetnames == EXPECTED_TABS


def test_non_accepted_cost_cells_read_excluded(book, result):
    sheet = book["BOM - Raw Materials"]
    unmatched = {line.name for line in result.ingredients if not line.match.accepted}
    assert unmatched
    seen = set()
    for row in sheet.iter_rows(min_row=4, values_only=True):
        if row[0] in unmatched:
            seen.add(row[0])
            assert row[20] == EXCLUDED_DISPLAY   # cost/bottle primary
            assert row[21] == EXCLUDED_DISPLAY   # low
            assert row[22] == EXCLUDED_DISPLAY   # high
    assert seen == unmatched


def test_bom_total_sums_accepted_lines_only(book, result):
    sheet = book["BOM - Raw Materials"]
    total = None
    for row in sheet.iter_rows(min_row=4, values_only=True):
        if row[0] and str(row[0]).startswith("TOTAL"):
            total = row[20]
    assert total == pytest.approx(result.summary.raw_materials)


def test_review_flags_tab_keeps_every_owner(book):
    sheet = book["Review Flags"]
    text = [row[0] for row in sheet.iter_rows(min_row=4, max_col=1, values_only=True)]
    for owner in ("Purchasing", "R&D", "Operations", "Sales", "Finance"):
        assert owner in text


def test_customer_tab_says_it_is_demonstration_data(book):
    """The shipped tables are illustrative, and the tab must not read as a draft
    of a real quote."""
    sheet = book["Customer Quote Summary"]
    banner = str(sheet.cell(row=2, column=1).value)
    assert "NOT A QUOTE" in banner and "DEMONSTRATION DATA" in banner


def test_customer_tab_is_marked_internal_draft_on_operational_data(
    reference, heart_health
):
    import copy

    catalogue = copy.deepcopy(reference)
    catalogue.rates = {**catalogue.rates, "dataset_is_demonstration": "0"}
    result = run_pipeline(heart_health, catalogue, as_of=AS_OF)
    sheet = load_workbook(io.BytesIO(build_workbook(result)))["Customer Quote Summary"]
    assert "INTERNAL DRAFT" in str(sheet.cell(row=2, column=1).value)


def test_missing_fields_appear_on_the_product_summary(reference, heart_health):
    heart_health.missing_fields = ["MOQ", "Timeline"]
    result = run_pipeline(heart_health, reference, as_of=AS_OF)
    sheet = load_workbook(io.BytesIO(build_workbook(result)))["Product Summary"]
    text = " ".join(str(row[0]) for row in sheet.iter_rows(max_col=1, values_only=True))
    assert "MOQ" in text and "Timeline" in text


def test_pdf_renders(result):
    data = build_quote_pdf(result)
    assert data.startswith(b"%PDF-")
    assert len(data) > 2000
