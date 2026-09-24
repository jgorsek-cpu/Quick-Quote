"""Operations' September 2026 data: crew, set up, cleaning, batches and lines.

Every number here came from Operations rather than from a placeholder, and
several of them changed answers the engine was already giving. The behaviour
each one drives is asserted here.
"""
import copy
from pathlib import Path

import pytest

from quickquote.engine.manufacturing import estimate_manufacturing
from quickquote.engine.pipeline import run_pipeline
from quickquote.schemas import FormulaLine, ParsedQuote, ProductSpec

from .conftest import AS_OF


def spec(form="Capsule", **kwargs):
    defaults = dict(dosage_form=form, capsule_size="0", capsule_type="Vegetable",
                    count_per_bottle=60, servings_per_bottle=30,
                    annual_volume_bottles=2000)
    defaults.update(kwargs)
    return ProductSpec(**defaults)


def quote(form="Capsule", lines=None, customer=None, **kwargs):
    if customer is not None:
        kwargs["customer"] = customer
    return ParsedQuote(
        product=spec(form, **kwargs),
        formula=lines or [FormulaLine("Ashwagandha Powder", 500.0)],
        packaging=[], source_name="test", source_format="manual",
    )


class TestCrewSize:
    """Labor is per operator. Packaging runs four of them."""

    def test_the_rate_multiplies_labor_but_not_overhead(self, reference):
        centre = reference.work_centres["Packaging"]
        assert centre.crew_size == 4
        assert centre.combined_rate == pytest.approx(
            centre.labor_rate_per_hour * 4 + centre.overhead_rate_per_hour)

    def test_a_one_operator_centre_is_unchanged(self, reference):
        centre = reference.work_centres["Tablet press"]
        assert centre.crew_size == 1
        assert centre.combined_rate == pytest.approx(centre.single_operator_rate)

    def test_the_packaging_line_overrides_the_centre_crew(self, reference):
        """PKG1 runs four operators, PKG2&3 five, on the same work centre."""
        four = reference.centre_rate("Packaging", 4)
        five = reference.centre_rate("Packaging", 5)
        labor = reference.work_centres["Packaging"].labor_rate_per_hour
        assert five - four == pytest.approx(labor)

    def test_it_reaches_the_bottling_step(self, reference):
        estimate, _ = estimate_manufacturing(spec(), 7, reference)
        bottling = next(step for step in estimate.steps if step.step == "Bottling")
        assert bottling.crew_size == 4
        assert bottling.rate_per_hour == pytest.approx(reference.centre_rate("Packaging", 4))

    def test_the_specification_mode_keeps_one_operator(self, spec_reference):
        """The published example predates crew size and must still reproduce."""
        estimate, _ = estimate_manufacturing(spec(), 7, spec_reference)
        bottling = next(step for step in estimate.steps if step.step == "Bottling")
        assert bottling.rate_per_hour == pytest.approx(spec_reference.combined_rate)


class TestSetUpTime:
    def test_set_up_is_added_to_the_run(self, reference):
        estimate, _ = estimate_manufacturing(spec(), 7, reference)
        encapsulation = next(step for step in estimate.steps if step.step == "Encapsulation")
        assert encapsulation.setup_hours == pytest.approx(2.0)
        assert encapsulation.hours == pytest.approx(
            encapsulation.run_hours + encapsulation.setup_hours)

    def test_it_is_paid_once_per_batch(self, reference):
        """Two batches means two set ups, not one spread thinner."""
        one = estimate_manufacturing(spec(), 7, reference, blend_kg_per_bottle=0.001)[0]
        many = estimate_manufacturing(spec(), 7, reference, blend_kg_per_bottle=10.0)[0]
        assert one.batches == 1 and many.batches > 1
        step_one = next(s for s in one.steps if s.step == "Compounding")
        step_many = next(s for s in many.steps if s.step == "Compounding")
        assert step_many.setup_hours == pytest.approx(step_one.setup_hours * many.batches)

    def test_a_cleaning_step_carries_no_set_up_of_its_own(self, reference):
        estimate, _ = estimate_manufacturing(spec(), 7, reference)
        cleaning = [step for step in estimate.steps if "cleaning" in step.step.lower()]
        assert cleaning and all(step.setup_hours is None for step in cleaning)

    def test_a_missing_set_up_time_understates_rather_than_blocks(self, reference):
        """One missing figure must not make a step that runs unquotable."""
        catalogue = copy.deepcopy(reference)
        catalogue.machines = [
            type(machine)(**{**machine.__dict__, "setup_hours": None})
            for machine in catalogue.machines
        ]
        estimate, _ = estimate_manufacturing(spec(annual_volume_bottles=500), 7, catalogue)
        assert estimate.machine == "Schaefer"
        assert estimate.estimated
        encapsulation = next(step for step in estimate.steps if step.step == "Encapsulation")
        assert encapsulation.setup_hours is None
        assert encapsulation.cost_per_bottle is not None

    def test_every_encapsulator_now_has_one(self, reference):
        """Operations closed the last gap in September 2026."""
        assert all(machine.setup_hours and machine.cleaning_hours
                   for machine in reference.machines)


class TestCleaning:
    def test_encapsulation_cleaning_follows_the_machine(self, reference):
        """Two hours on the 705, four on the 1505, six on the 3005."""
        hours = {}
        for volume, machine in ((3000, "BOSCH 705"), (10000, "BOSCH 1505"),
                                (40000, "BOSCH 3005")):
            estimate, _ = estimate_manufacturing(
                spec(annual_volume_bottles=volume), 7, reference)
            assert estimate.machine == machine
            step = next(s for s in estimate.steps if s.step == "Encapsulation cleaning")
            hours[machine] = step.run_hours
        assert hours == {"BOSCH 705": 2.0, "BOSCH 1505": 4.0, "BOSCH 3005": 6.0}

    def test_cleaning_is_paid_once_per_batch(self, reference):
        many = estimate_manufacturing(spec(), 7, reference, blend_kg_per_bottle=10.0)[0]
        step = next(s for s in many.steps if s.step == "Blend cleaning")
        assert step.per_batch
        assert step.run_hours == pytest.approx(0.5 * many.batches)


class TestBatches:
    """Blender capacity is a volume limit, so density decides the weight."""

    def test_a_denser_blend_fits_in_fewer_batches(self, reference):
        loose = reference.batches_for(5000, 0.4)
        dense = reference.batches_for(5000, 0.8)
        assert loose[0] > dense[0]

    def test_a_small_run_uses_a_small_blender(self, reference):
        _, blender, _ = reference.batches_for(50, 0.4)
        assert blender.blender == "300 L"

    def test_a_run_too_large_for_any_blender_is_split(self, reference):
        batches, blender, _ = reference.batches_for(10_000, 0.4)
        assert blender.blender == "8500 L"
        assert batches == 4          # 10,000 kg / 2,720 kg

    def test_the_default_density_is_operations_quoting_assumption(self, reference):
        assert reference.default_blend_density == pytest.approx(0.4)
        assert reference.batches_for(1000)[2] == pytest.approx(0.4)

    def test_per_batch_cost_does_not_amortise_away(self, reference):
        """The flaw a single-batch model hides: at volume, compounding used to
        approach zero per bottle. It cannot, because the batches repeat."""
        small = estimate_manufacturing(
            spec(annual_volume_bottles=2000), 7, reference, blend_kg_per_bottle=0.5)[0]
        large = estimate_manufacturing(
            spec(annual_volume_bottles=2_000_000), 7, reference, blend_kg_per_bottle=0.5)[0]
        assert large.batches > small.batches
        small_cost = next(s for s in small.steps if s.step == "Compounding").cost_per_bottle
        large_cost = next(s for s in large.steps if s.step == "Compounding").cost_per_bottle
        # A thousandfold volume increase, but compounding per bottle settles
        # rather than vanishing.
        assert large_cost > small_cost / 10

    def test_the_blend_density_comes_from_the_formula_when_it_is_known(self, reference):
        result = run_pipeline(quote(), reference, as_of=AS_OF)
        assert not result.manufacturing.blend_density_assumed
        assert result.manufacturing.blend_density_g_ml > 0

    def test_one_unknown_density_falls_back_to_the_assumption(self, reference):
        result = run_pipeline(
            quote(lines=[FormulaLine("Ashwagandha Powder", 500.0),
                         FormulaLine("Wholly Unknown Material", 100.0)]),
            reference, as_of=AS_OF)
        assert result.manufacturing.blend_density_assumed
        assert result.manufacturing.blend_density_g_ml == pytest.approx(0.4)


class TestPackagingLines:
    def test_a_count_between_rows_takes_the_slower(self, reference):
        assert reference.packaging_line_row(75, "PKG1").bottles_per_hour == 1620

    def test_a_count_past_the_last_row_takes_the_last(self, reference):
        assert reference.packaging_line_row(5000, "PKG1").count_per_bottle == 200

    def test_the_assumed_line_is_named_with_what_the_other_would_cost(self, reference):
        result = run_pipeline(quote(), reference, as_of=AS_OF)
        line = [flag for flag in result.flags if flag.item == "Packaging line"]
        assert line, "the assumed line must be visible on the quote"
        assert "PKG1" in line[0].reason and "PKG2&3" in line[0].reason
        assert "operators" in line[0].reason


class TestFormsOperationsUnblocked:
    @pytest.mark.parametrize("form", ["Powder", "Packet"])
    def test_it_can_now_be_quoted(self, reference, form):
        result = run_pipeline(quote(form), reference, as_of=AS_OF)
        assert result.manufacturing.estimated
        assert result.summary.manufacturing > 0

    def test_a_line_rated_in_bottles_fills_bottles_not_capsules(self, reference):
        """The powder line runs 1,200 bottles an hour whatever goes in them."""
        estimate, _ = estimate_manufacturing(
            spec("Powder", count_per_bottle=60, annual_volume_bottles=2000), 7, reference)
        fill = next(step for step in estimate.steps if step.step == "Powder fill")
        assert fill.run_hours == pytest.approx(2000 / 1200)

    def test_a_tablet_is_still_blocked_by_granulation(self, reference):
        """The press rate arrived; the Chilsinator's did not."""
        estimate, _ = estimate_manufacturing(spec("Tablet"), 7, reference)
        assert not estimate.estimated
        assert estimate.uncosted_critical_steps == ["Granulation"]

    def test_a_gummy_is_still_blocked(self, reference):
        """Finance confirmed the gummy cost model does not exist yet."""
        estimate, _ = estimate_manufacturing(spec("Gummy"), 7, reference)
        assert not estimate.estimated


class TestConfirmedRates:
    @pytest.mark.parametrize("centre,speed", [
        ("Schaefer", 8000), ("705", 32000), ("1505", 75000), ("3005", 110000),
        ("Tablet press", 141000), ("Powder Fill", 1200), ("Pakrapid", 3500),
    ])
    def test_operations_speeds_are_on_file_and_marked_confirmed(
        self, reference, centre, speed
    ):
        rate = reference.run_rates[centre]
        assert rate.units_per_hour == speed and rate.confirmed

    def test_a_rate_operations_did_not_supply_is_not_marked_confirmed(self, reference):
        """Granulation is the one throughput still missing."""
        unconfirmed = [name for name, rate in reference.run_rates.items()
                       if not rate.confirmed]
        assert unconfirmed == ["Chilsinator"]

    def test_the_schaefer_rate_arrived(self, reference):
        rate = reference.run_rates["Schaefer"]
        assert rate.units_per_hour == 8000 and rate.confirmed

    def test_the_encapsulation_bands_match_what_operations_confirmed(self, reference):
        bands = {m.machine: (m.min_capsules, m.max_capsules) for m in reference.machines}
        assert bands["Schaefer"] == (0, 100_000)
        assert bands["BOSCH 705"] == (100_000, 400_000)
        assert bands["BOSCH 1505"] == (400_000, 2_000_000)
        assert bands["BOSCH 3005"] == (2_000_000, None)


class TestCompoundingDisagreement:
    def test_operations_own_timings_are_available(self, reference):
        """Five minutes an ingredient, 45 minutes of blending, 10s per 20 kg."""
        hours = reference.compounding_hours_from_operations(9, 3600)
        assert hours == pytest.approx((5 * 9 + 45) / 60 + (3600 / 20) * 10 / 3600)

    def test_the_gap_against_the_bands_is_flagged_not_resolved(self, reference):
        result = run_pipeline(quote(), reference, as_of=AS_OF)
        gap = [flag for flag in result.flags if flag.item == "Compounding time"]
        assert gap and gap[0].owner == "Operations"
        assert "component-count bands" in gap[0].reason

    def test_the_engine_still_charges_the_bands(self, reference):
        """Flagging a disagreement is not the same as picking a side."""
        estimate, _ = estimate_manufacturing(spec(), 7, reference)
        step = next(s for s in estimate.steps if s.step == "Compounding")
        assert step.run_hours == pytest.approx(reference.compounding_hours(7))


class TestUnknownDensityIsVisible:
    def test_an_undeterminable_class_does_not_borrow_the_default_row(self, reference):
        """"default" names a class that could not be worked out, not a class."""
        known, was_defaulted = reference.density_for("ashwagandha powder", "botanical_powder")
        assert not was_defaulted and known > 0
        _, unknown_defaulted = reference.density_for(None, "default")
        assert unknown_defaulted

    def test_it_reaches_r_and_d_as_a_flag(self, reference):
        result = run_pipeline(
            quote(lines=[FormulaLine("Wholly Unknown Material", 100.0)]),
            reference, as_of=AS_OF)
        assert any("Bulk density not on file" in flag.reason for flag in result.flags)


class TestDatasetProvenance:
    """A quote costed against demonstration tables is complete, confident and
    fiction. Its part codes look exactly like real ones, so it has to say so."""

    def test_the_shipped_tables_declare_themselves_a_demonstration(self, reference):
        assert reference.is_demonstration

    def test_every_quote_on_them_carries_a_blocking_flag(self, reference):
        result = run_pipeline(quote(), reference, as_of=AS_OF)
        notice = [flag for flag in result.flags if flag.item == "Reference data"]
        assert notice and notice[0].severity == "blocking"
        assert "NOT FOR QUOTING" in notice[0].reason
        assert result.dataset_is_demonstration

    def test_operational_tables_raise_no_such_flag(self, reference):
        catalogue = copy.deepcopy(reference)
        catalogue.rates = {**catalogue.rates, "dataset_is_demonstration": "0",
                           "dataset_label": "DrVita operational data loaded 2026-09-24"}
        result = run_pipeline(quote(), catalogue, as_of=AS_OF)
        assert not result.dataset_is_demonstration
        assert not [flag for flag in result.flags if flag.item == "Reference data"]
        assert result.dataset_label.startswith("DrVita operational")

    def test_the_loader_script_clears_the_marker(self):
        """Otherwise seeding the real directory would stamp it a demonstration."""
        script = (Path(__file__).resolve().parents[1]
                  / "scripts" / "load_real_data.sh").read_text()
        assert "dataset_is_demonstration" in script
        assert script.index("cp backend/quickquote/reference/data/*.csv") < \
               script.index("dataset_is_demonstration")


class TestMarginRules:
    """Finance's requirements: by account, on two cost bases, with a form floor."""

    @pytest.mark.parametrize("customer,rule", [
        ("Costco Wholesale", "costco"),
        ("WALMART STORES INC", "walmart"),
        ("Nature's Lab", "natures_lab"),
        ("Natures Lab LLC", "natures_lab"),
        ("Vitamin Shoppe", "default"),
        (None, "default"),
    ])
    def test_the_account_decides_the_rule(self, reference, customer, rule):
        assert reference.margin_rule_for(customer).rule == rule

    def test_a_near_miss_does_not_match(self, reference):
        """"Costa Coffee" is not Costco."""
        assert reference.margin_rule_for("Costa Coffee").rule == "default"

    def test_costco_and_walmart_are_held_higher_than_the_rest(self, reference):
        costco = reference.margin_rule_for("Costco Wholesale")
        standard = reference.margin_rule_for("Vitamin Shoppe")
        assert costco.min_margin_incl_oh_pct == 30.0
        assert standard.min_margin_incl_oh_pct == 20.0

    def test_the_standard_terms_carry_two_tests(self, reference):
        standard = reference.margin_rule_for("Vitamin Shoppe")
        assert standard.min_margin_incl_oh_pct == 20.0
        assert standard.min_margin_excl_oh_pct == 30.0

    def test_both_tests_produce_a_price_and_one_binds(self, reference):
        result = run_pipeline(quote(customer="Vitamin Shoppe"), reference, as_of=AS_OF)
        assert len(result.pricing) == 2
        binding = [item for item in result.pricing if item.binding]
        assert len(binding) == 1
        assert binding[0].price_per_bottle == max(
            item.price_per_bottle for item in result.pricing)

    def test_the_excluding_overhead_price_uses_the_smaller_cost(self, reference):
        result = run_pipeline(quote(customer="Vitamin Shoppe"), reference, as_of=AS_OF)
        excl = next(i for i in result.pricing if i.cost_basis == "excluding overhead")
        incl = next(i for i in result.pricing if i.cost_basis == "including overhead")
        assert excl.cost_per_bottle < incl.cost_per_bottle
        assert incl.cost_per_bottle == pytest.approx(result.summary.primary_per_bottle)
        assert excl.cost_per_bottle == pytest.approx(
            result.summary.cost_excluding_overhead)

    def test_overhead_is_what_separates_the_two_bases(self, reference):
        result = run_pipeline(quote(customer="Vitamin Shoppe"), reference, as_of=AS_OF)
        summary = result.summary
        assert summary.manufacturing == pytest.approx(
            summary.manufacturing_labor + summary.manufacturing_overhead)
        assert summary.cost_excluding_overhead == pytest.approx(
            summary.primary_per_bottle - summary.manufacturing_overhead)

    def test_a_named_account_is_priced_higher(self, reference):
        cheap = run_pipeline(quote(customer="Vitamin Shoppe"), reference, as_of=AS_OF)
        rich = run_pipeline(quote(customer="Nature's Lab"), reference, as_of=AS_OF)
        assert max(i.price_per_bottle for i in rich.pricing) > \
               max(i.price_per_bottle for i in cheap.pricing)


class TestDosageFormFloor:
    """A tablet, powder or stick pack may go to 20% whichever account it is."""

    @pytest.mark.parametrize("form,expected", [
        ("Tablet", "tablet"), ("Chewable Tablet", "tablet"),
        ("Powder", "powder"), ("Stick Pack", "stick pack"),
        ("Capsule", None), ("Softgel", None), ("Gummy", None),
    ])
    def test_which_forms_qualify(self, reference, form, expected):
        assert reference.form_allows_floor(form) == expected

    def test_it_relaxes_a_named_account(self, reference):
        capsule = run_pipeline(quote(customer="Costco Wholesale"), reference, as_of=AS_OF)
        powder = run_pipeline(quote(customer="Costco Wholesale", form="Powder"),
                              reference, as_of=AS_OF)
        assert [i.target_margin_pct for i in capsule.pricing] == [30.0]
        assert [i.target_margin_pct for i in powder.pricing] == [20.0]
        assert powder.pricing[0].form_floor_applied
        assert "relaxed from 30%" in powder.pricing[0].basis

    def test_it_never_tightens_a_requirement(self, reference):
        """The standard terms are already at the floor; the form changes nothing."""
        result = run_pipeline(quote(customer="Vitamin Shoppe", form="Powder"),
                              reference, as_of=AS_OF)
        incl = next(i for i in result.pricing if i.cost_basis == "including overhead")
        assert incl.target_margin_pct == 20.0
        assert not incl.form_floor_applied

    def test_a_capsule_keeps_the_full_requirement(self, reference):
        result = run_pipeline(quote(customer="Nature's Lab"), reference, as_of=AS_OF)
        assert result.pricing[0].target_margin_pct == 65.0
        assert not result.pricing[0].form_floor_applied


class TestUnconfirmedMargins:
    def test_an_unconfirmed_rule_is_flagged_to_finance(self, reference):
        catalogue = copy.deepcopy(reference)
        catalogue.margin_rules = [
            type(rule)(**{**rule.__dict__, "confirmed": False})
            for rule in catalogue.margin_rules
        ]
        result = run_pipeline(quote(customer="Costco Wholesale"), catalogue, as_of=AS_OF)
        margin = [flag for flag in result.flags if flag.item == "Margin targets"]
        assert margin and margin[0].severity == "blocking"

    def test_a_confirmed_rule_raises_nothing(self, reference):
        result = run_pipeline(quote(customer="Costco Wholesale"), reference, as_of=AS_OF)
        assert not [flag for flag in result.flags if flag.item == "Margin targets"]

    def test_a_quote_with_no_customer_says_which_terms_it_used(self, reference):
        result = run_pipeline(quote(), reference, as_of=AS_OF)
        margin = [flag for flag in result.flags if flag.item == "Margin targets"]
        assert margin and "Walmart" in margin[0].reason

    def test_the_price_is_still_offered_with_the_caveat(self, reference):
        """Flagging it is not the same as withholding it."""
        result = run_pipeline(quote(), reference, as_of=AS_OF)
        assert result.pricing and all(item.price_per_bottle > 0 for item in result.pricing)


class TestDescriptionEncoding:
    """Global Shop exports cp1252 text that arrives decoded as cp437."""

    @pytest.mark.parametrize("broken,fixed", [
        ("KSM-66« ASHWAGANDHA EXTRACT", "KSM-66® ASHWAGANDHA EXTRACT"),
        ("3.25 x 3.5 WOMENÆS WHITE FILM", "3.25 x 3.5 WOMEN’S WHITE FILM"),
        ("5-HTP, 98% (HTPurityÖ)", "5-HTP, 98% (HTPurity™)"),
        ("32oz WHITE HDPE ôLIPö", "32oz WHITE HDPE “LIP”"),
    ])
    def test_it_is_repaired(self, broken, fixed):
        from quickquote.reference.ingest import repair_mojibake

        assert repair_mojibake(broken) == fixed

    @pytest.mark.parametrize("text", [
        "Plain ASCII Description", "Café Latte Flavour", "Açaí Berry Extract",
        "PURÉE BASE", "Jalapeño Powder",
    ])
    def test_a_real_accent_is_left_alone(self, text):
        """The repair must not mangle a description that was always correct."""
        from quickquote.reference.ingest import repair_mojibake

        assert repair_mojibake(text) == text


class TestPoExportColumns:
    """The full exports carry description, vendor and PO number."""

    def _book(self, tmp_path, rows):
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        for row in rows:
            sheet.append(list(row))
        path = tmp_path / "po.xlsx"
        book.save(path)
        return path

    HEADER = ("Purchase Order Date", "Part Number", "Description",
              "Purchase Order Number", "Vendor", "Vendor Name", "Order Qty", "Unit Cost")

    def test_vendors_are_read_and_counted(self, tmp_path):
        from quickquote.reference.ingest import IngestReport, aggregate, read_po_transactions

        path = self._book(tmp_path, [
            self.HEADER,
            ("2024-01-01", "RNMSM1", "MSM", "1", "FUL", "FULLER ENTERPRISE", 100, 10.0),
            ("2025-01-01", "RNMSM1", "MSM", "2", "FUL", "FULLER ENTERPRISE", 100, 11.0),
            ("2026-01-01", "RBASH6", "ASHWAGANDHA", "3", "SKP", "SHRI KARTIKEYA", 50, 20.0),
            ("2025-06-01", "RBASH6", "ASHWAGANDHA", "4", "ASI", "A.S.I. INTERNATIONAL", 50, 22.0),
        ])
        report = IngestReport()
        rows = {r["part_number"]: r
                for r in aggregate(read_po_transactions([path], report), {}, report)}
        assert report.vendors_present
        assert rows["RNMSM1"]["unique_vendor_count"] == "1"
        assert rows["RNMSM1"]["latest_vendor"] == "FULLER ENTERPRISE"
        assert rows["RBASH6"]["unique_vendor_count"] == "2"
        # Most recent PO first, so the vendor named is the one to call.
        assert rows["RBASH6"]["latest_vendor"] == "SHRI KARTIKEYA"

    def test_a_description_missing_from_the_master_comes_from_the_po(self, tmp_path):
        from quickquote.reference.ingest import IngestReport, aggregate, read_po_transactions

        path = self._book(tmp_path, [
            self.HEADER,
            ("2026-01-01", "RNEW1", "BRAND NEW MATERIAL", "9", "V", "VENDOR", 10, 5.0),
        ])
        report = IngestReport()
        rows = aggregate(read_po_transactions([path], report), {}, report)
        assert rows[0]["description"] == "BRAND NEW MATERIAL"
        assert report.descriptions_from_po == 1
        assert report.parts_without_description == []

    def test_the_item_master_still_wins(self, tmp_path):
        from quickquote.reference.ingest import IngestReport, aggregate, read_po_transactions

        path = self._book(tmp_path, [
            self.HEADER,
            ("2026-01-01", "RNMSM1", "po spelling", "9", "V", "VENDOR", 10, 5.0),
        ])
        report = IngestReport()
        master = {"RNMSM1": {"description": "master spelling", "uom": "KG"}}
        rows = aggregate(read_po_transactions([path], report), master, report)
        assert rows[0]["description"] == "master spelling"

    def test_a_po_description_is_repaired_too(self, tmp_path):
        from quickquote.reference.ingest import IngestReport, aggregate, read_po_transactions

        path = self._book(tmp_path, [
            self.HEADER,
            ("2026-01-01", "RBASH6", "KSM-66« ASHWAGANDHA", "9", "V", "VENDOR", 10, 5.0),
        ])
        report = IngestReport()
        rows = aggregate(read_po_transactions([path], report), {}, report)
        assert rows[0]["description"] == "KSM-66® ASHWAGANDHA"

    def test_a_single_supplier_part_reaches_purchasing(self, reference):
        """The flag was already written; it was the vendor column that was missing."""
        import copy as _copy
        from datetime import date as _date
        from quickquote.reference.loader import PoRow

        catalogue = _copy.deepcopy(reference)
        catalogue.identity_index = None
        catalogue.po_rows = [
            PoRow(part_number="RNMSM1", description="MSM", uom="KG",
                  latest_unit_cost=11.0, min_unit_cost_ever=11.0, max_unit_cost_ever=11.0,
                  latest_po_date=_date(2026, 8, 1), latest_vendor="FULLER ENTERPRISE",
                  unique_vendor_count=1, po_count=7, total_spend=1000.0),
        ]
        result = run_pipeline(
            quote(lines=[FormulaLine("MSM", 500.0)]), catalogue, as_of=AS_OF)
        sole = [f for f in result.flags if "Single-supplier" in f.reason]
        assert sole and "FULLER ENTERPRISE" in sole[0].reason


class TestUnitOfMeasure:
    """Global Shop writes a thousand three ways, and reading one as "each"
    costs a thousand times too much."""

    @pytest.mark.parametrize("purchasing,stock,expected", [
        ("EA", "EA", "EA"),
        ("KG", "KG", "KG"),
        ("M", "M", "M"),
        ("K", "K", "M"),        # most of DrVita's catalogue writes K
        ("KM", "M", "M"),       # bought by the roll, stocked by the thousand
        ("RL", "EA", "EA"),     # a roll is the each it is bought as
        ("", "KG", "KG"),       # purchasing blank, stocking usable
    ])
    def test_it_resolves_to_a_unit_the_engine_can_price(self, purchasing, stock, expected):
        from quickquote.reference.ingest import normalise_uom

        assert normalise_uom(purchasing, stock) == expected

    def test_a_unit_it_cannot_price_is_refused_not_guessed(self):
        from quickquote.reference.ingest import normalise_uom

        assert normalise_uom("CS", "IN") is None

    def test_the_purchasing_unit_wins_where_both_are_usable(self):
        from quickquote.reference.ingest import normalise_uom

        assert normalise_uom("KG", "EA") == "KG"

    def test_a_per_thousand_part_is_not_costed_as_each(self, tmp_path):
        """A 67mm shrink band is $76.80 a thousand, not $76.80 a bottle."""
        from openpyxl import Workbook
        from quickquote.reference.ingest import IngestReport, aggregate, read_po_transactions

        book = Workbook()
        sheet = book.active
        sheet.append(["Purchase Order Date", "Part Number", "Description",
                      "Purchase Order Number", "Vendor", "Vendor Name",
                      "Order Qty", "Unit Cost"])
        sheet.append(["2026-04-14", "KSBR67NP", "67MM CLEAR SHRINK BAND", "1",
                      "AMS", "AMERI-SEAL", 5, 76.8])
        path = tmp_path / "po.xlsx"
        book.save(path)

        report = IngestReport()
        master = {"KSBR67NP": {"description": "67MM CLEAR SHRINK BAND",
                               "uom": "KM", "stock_uom": "M"}}
        row = aggregate(read_po_transactions([path], report), master, report)[0]
        assert row["uom"] == "M"
        assert row["uom_basis"] == "purchasing unit"
        assert report.uom_inferred == []

    def test_a_guessed_unit_is_recorded(self, tmp_path):
        from openpyxl import Workbook
        from quickquote.reference.ingest import IngestReport, aggregate, read_po_transactions

        book = Workbook()
        sheet = book.active
        sheet.append(["Purchase Order Date", "Part Number", "Order Qty", "Unit Cost"])
        sheet.append(["2026-04-14", "KFC16", 2, 80.1])
        path = tmp_path / "po.xlsx"
        book.save(path)

        report = IngestReport()
        master = {"KFC16": {"description": "16 GRAM COIL RAYON 18lbs/CS",
                            "uom": "CS", "stock_uom": "IN"}}
        row = aggregate(read_po_transactions([path], report), master, report)[0]
        assert row["uom_basis"] == "inferred"
        assert report.uom_inferred == ["KFC16"]


class TestPackOutCompletion:
    def test_naming_one_component_does_not_lose_the_others(self, reference):
        """A rep who lists a bottle must not silently lose the capsule shell."""
        from quickquote.schemas import PackagingLine

        parsed = quote(count_per_bottle=60, servings_per_bottle=30)
        parsed.packaging = [PackagingLine("bottle", "150CC WHITE HDPE PACKER 38-400")]
        result = run_pipeline(parsed, reference, as_of=AS_OF)
        roles = {line.role for line in result.packaging}
        assert "capsule_shell" in roles and "bottle" in roles

    def test_a_named_component_is_left_alone(self, reference):
        from quickquote.schemas import PackagingLine

        parsed = quote(count_per_bottle=60, servings_per_bottle=30)
        parsed.packaging = [PackagingLine("bottle", "400CC WHITE HDPE BOTTLE 45/400")]
        result = run_pipeline(parsed, reference, as_of=AS_OF)
        bottles = [line for line in result.packaging if line.role == "bottle"]
        assert len(bottles) == 1 and "400" in bottles[0].description


class TestBottleSizing:
    def test_the_cheapest_bottle_that_fits_wins_not_the_smallest(self, reference):
        """A 100cc packer costs more than the 150cc that DrVita buy in volume."""
        import copy as _copy
        from datetime import date as _date
        from quickquote.engine.derive import choose_bottle
        from quickquote.reference.loader import PoRow

        catalogue = _copy.deepcopy(reference)
        catalogue.identity_index = None
        catalogue.po_rows = [
            PoRow(part_number="KTTP100", description="100cc WHITE HDPE PACKER 38-400",
                  uom="EA", latest_unit_cost=0.15, min_unit_cost_ever=0.15,
                  max_unit_cost_ever=0.15, latest_po_date=_date(2026, 4, 1),
                  latest_vendor="V", unique_vendor_count=1, po_count=3, total_spend=1.0),
            PoRow(part_number="KTTP150", description="150CC WHITE HDPE PACKER 38-400",
                  uom="EA", latest_unit_cost=0.119, min_unit_cost_ever=0.119,
                  max_unit_cost_ever=0.119, latest_po_date=_date(2026, 4, 1),
                  latest_vendor="V", unique_vendor_count=1, po_count=3, total_spend=1.0),
        ]
        choice = choose_bottle(spec(capsule_size="00", count_per_bottle=60,
                                    servings_per_bottle=30), catalogue)
        assert choice.row.part_number == "KTTP150"
        assert "costs more" in choice.basis

    def test_drvitas_own_wording_reaches_the_bottle_identity(self, reference):
        """The catalogue says "PACKER" where the identity said "bottle"."""
        from quickquote.engine.identity import resolve_identity

        resolved = resolve_identity("150CC WHITE HDPE PACKER 38-400", reference,
                                    packaging=True)
        assert resolved is not None and resolved.canonical == "hdpe bottle white"


class TestBottlingCleaning:
    def test_every_bottled_form_cleans_its_line(self, reference):
        """Operations supplied the hour; no route was using it."""
        for form in ("capsule", "tablet", "softgel", "gummy"):
            steps = [step.step for step in reference.route_for(form)]
            assert "Bottling cleaning" in steps, form

    def test_it_is_costed_at_the_packaging_crew(self, reference):
        estimate, _ = estimate_manufacturing(spec(count_per_bottle=60), 7, reference)
        step = next(s for s in estimate.steps if s.step == "Bottling cleaning")
        assert step.crew_size == 4 and step.run_hours == pytest.approx(1.0)
        assert step.cost_per_bottle > 0
