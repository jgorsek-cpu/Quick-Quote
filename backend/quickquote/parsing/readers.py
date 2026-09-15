"""Format detection and reading.

Every supported format is reduced to the same intermediate shape: a list of
tables plus a list of text lines. Nothing is interpreted here.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Table:
    """A rectangular block of cells read from one sheet, page or table."""

    rows: list[list[str]] = field(default_factory=list)
    source: str = ""


@dataclass
class Document:
    """The raw, uninterpreted reading of an input file."""

    tables: list[Table] = field(default_factory=list)
    lines: list[str] = field(default_factory=list)
    source_name: str = ""
    source_format: str = ""
    notes: list[str] = field(default_factory=list)


class UnsupportedFormat(ValueError):
    """Raised when a file extension has no parser."""


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rows_to_lines(rows: list[list[str]]) -> list[str]:
    lines = []
    for row in rows:
        text = " | ".join(cell for cell in row if cell)
        if text.strip():
            lines.append(text)
    return lines


def read_xlsx(data: bytes, name: str) -> Document:
    from openpyxl import load_workbook

    document = Document(source_name=name, source_format="xlsx")
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = [[_cell(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            rows = [row for row in rows if any(cell for cell in row)]
            if not rows:
                continue
            document.tables.append(Table(rows=rows, source=sheet.title))
            document.lines.extend(_rows_to_lines(rows))
    finally:
        workbook.close()
    return document


def read_pdf(data: bytes, name: str) -> Document:
    import pdfplumber

    document = Document(source_name=name, source_format="pdf")
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for index, page in enumerate(pdf.pages, start=1):
            for table in page.extract_tables() or []:
                rows = [[_cell(cell) for cell in row] for row in table]
                rows = [row for row in rows if any(cell for cell in row)]
                if rows:
                    document.tables.append(Table(rows=rows, source=f"page {index}"))
            text = page.extract_text() or ""
            document.lines.extend(
                line.strip() for line in text.splitlines() if line.strip()
            )
    if not document.lines and not document.tables:
        document.notes.append(
            "No extractable text found - the PDF may be a scan. "
            "Fields could not be read and are recorded as missing."
        )
    return document


def read_docx(data: bytes, name: str) -> Document:
    import docx

    document = Document(source_name=name, source_format="docx")
    source = docx.Document(io.BytesIO(data))
    for index, table in enumerate(source.tables, start=1):
        rows = [[_cell(cell.text) for cell in row.cells] for row in table.rows]
        rows = [row for row in rows if any(cell for cell in row)]
        if rows:
            document.tables.append(Table(rows=rows, source=f"table {index}"))
            document.lines.extend(_rows_to_lines(rows))
    for paragraph in source.paragraphs:
        text = paragraph.text.strip()
        if text:
            document.lines.append(text)
    return document


def read_text(data: bytes, name: str, source_format: str = "txt") -> Document:
    document = Document(source_name=name, source_format=source_format)
    text = data.decode("utf-8", errors="replace")
    raw_lines = [line.rstrip() for line in text.splitlines()]
    document.lines = [line for line in raw_lines if line.strip()]

    # Blank lines are kept for table detection: in a text quote sheet they are
    # what separates the ingredient block from the packaging block.
    table = _detect_table(raw_lines, name)
    if table is not None:
        document.tables.append(table)
    return document


def _detect_table(lines: list[str], name: str) -> Table | None:
    """Find the delimiter that yields the most multi-column rows.

    A quote sheet saved as text usually mixes ``Label: value`` lines with a
    delimited ingredient block. Sniffing the whole file fails on that mix, so
    each candidate delimiter is scored directly.
    """
    best_rows: list[list[str]] | None = None
    best_score = 0

    for delimiter in (",", "\t", "|", ";"):
        try:
            parsed = list(csv.reader(lines, delimiter=delimiter))
        except csv.Error:
            continue
        rows = [[cell.strip() for cell in row] for row in parsed]
        score = sum(1 for row in rows if len([c for c in row if c]) >= 3)
        if score > best_score:
            best_rows, best_score = rows, score

    if best_rows is None or best_score < 2:
        return None
    if not any(any(cell for cell in row) for row in best_rows):
        return None
    return Table(rows=best_rows, source=name)


_READERS = {
    ".xlsx": read_xlsx,
    ".xlsm": read_xlsx,
    ".pdf": read_pdf,
    ".docx": read_docx,
}


def read_document(data: bytes, filename: str) -> Document:
    """Detect the format from the file name and route to the right reader."""
    suffix = Path(filename).suffix.lower()
    if suffix in _READERS:
        return _READERS[suffix](data, filename)
    if suffix == ".csv":
        return read_text(data, filename, source_format="csv")
    if suffix in (".txt", ".text", ""):
        return read_text(data, filename, source_format="txt")
    raise UnsupportedFormat(
        f"'{suffix or filename}' is not a supported quote document. "
        "Supported formats: xlsx, xlsm, pdf, docx, txt, csv."
    )
