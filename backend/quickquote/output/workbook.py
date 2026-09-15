"""Six-tab Excel workbook generation.

Every quote produces the same six tabs with the same fields. Cost columns for
non-Accepted lines display "Excluded", never an estimate.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..config import ACCEPTED, EXCLUDED_DISPLAY, FLAG_OWNERS, HIGH, LOW, MEDIUM, NEEDS_REVIEW
from ..schemas import CostedIngredient, CostedPackaging, QuoteResult

NAVY = "1F3864"
LIGHT = "D9E2F3"
RED = "C00000"
RED_FILL = "FCE4E4"
AMBER_FILL = "FFF2CC"
GREEN_FILL = "E2EFDA"
GREY_FILL = "F2F2F2"

TITLE_FONT = Font(bold=True, size=14, color="FFFFFF")
HEAD_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
SMALL = Font(size=9, color="595959")

MONEY = '"$"#,##0.0000'
MONEY2 = '"$"#,##0.00'
PCT = "0.0%"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_FILL = {
    ACCEPTED: PatternFill("solid", fgColor=GREEN_FILL),
    NEEDS_REVIEW: PatternFill("solid", fgColor=AMBER_FILL),
}
CONFIDENCE_FILL = {
    HIGH: PatternFill("solid", fgColor=GREEN_FILL),
    MEDIUM: PatternFill("solid", fgColor=AMBER_FILL),
    LOW: PatternFill("solid", fgColor=RED_FILL),
}


def _title(sheet: Worksheet, text: str, width: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(width, 2))
    cell = sheet.cell(row=1, column=1, value=text)
    cell.font = TITLE_FONT
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    sheet.row_dimensions[1].height = 24


def _header_row(sheet: Worksheet, row: int, headers: list[str]) -> None:
    for column, text in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=column, value=text)
        cell.font = HEAD_FONT
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    sheet.row_dimensions[row].height = 30
    sheet.freeze_panes = sheet.cell(row=row + 1, column=1)


def _section(sheet: Worksheet, row: int, text: str, width: int = 6) -> int:
    cell = sheet.cell(row=row, column=1, value=text)
    cell.font = BOLD
    cell.fill = PatternFill("solid", fgColor=LIGHT)
    for column in range(2, width + 1):
        sheet.cell(row=row, column=column).fill = PatternFill("solid", fgColor=LIGHT)
    return row + 1


def _kv(sheet: Worksheet, row: int, label: str, value, number_format: str | None = None) -> int:
    sheet.cell(row=row, column=1, value=label).font = BOLD
    cell = sheet.cell(row=row, column=2, value=value)
    if number_format:
        cell.number_format = number_format
    return row + 1


def _widths(sheet: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _money_or_excluded(line: CostedIngredient | CostedPackaging, value: float | None):
    """Accepted lines show a number; everything else shows 'Excluded'."""
    if line.match.status != ACCEPTED or value is None:
        return EXCLUDED_DISPLAY
    return value


def _dash(value):
    return value if value not in (None, "") else "-"


# ------------------------------------------------------------- Tab 1

def _product_summary(book: Workbook, result: QuoteResult) -> None:
    sheet = book.create_sheet("Product Summary")
    _widths(sheet, [30, 46, 20, 20])
    _title(sheet, "PRODUCT SUMMARY", 4)

    product = result.product
    row = 3
    row = _section(sheet, row, "Identification", 4)
    row = _kv(sheet, row, "Customer", _dash(product.customer))
    row = _kv(sheet, row, "Brand", _dash(product.brand))
    row = _kv(sheet, row, "Formula", _dash(product.formula_name))
    row = _kv(sheet, row, "Customer SKU", _dash(product.customer_sku))
    row += 1

    row = _section(sheet, row, "Product form", 4)
    row = _kv(sheet, row, "Dosage form", _dash(product.dosage_form))
    row = _kv(sheet, row, "Capsule size", _dash(product.capsule_size))
    row = _kv(sheet, row, "Capsule type", _dash(product.capsule_type))
    row = _kv(sheet, row, "Serving size", _dash(product.serving_size))
    row = _kv(sheet, row, "Servings per bottle", _dash(product.servings_per_bottle))
    row = _kv(sheet, row, "Count per bottle", _dash(product.count_per_bottle))
    row += 1

    row = _section(sheet, row, "Commercial terms", 4)
    row = _kv(sheet, row, "Volume (bottles)", _dash(product.annual_volume_bottles))
    row = _kv(sheet, row, "MOQ", _dash(product.moq))
    row = _kv(sheet, row, "Timeline", _dash(product.timeline))
    row += 1

    row = _section(sheet, row, "Manufacturing assumptions", 4)
    row = _kv(sheet, row, "Machine", result.manufacturing.machine +
              (" (assumed)" if result.manufacturing.machine_assumed else ""))
    row = _kv(sheet, row, "Mfg loss factor",
              product.mfg_loss_factor if product.mfg_loss_factor else "1.05 (default)")
    row += 1

    row = _section(sheet, row, "Claims and testing", 4)
    row = _kv(sheet, row, "Claims", "; ".join(product.claims) if product.claims else "-")
    row = _kv(sheet, row, "Testing", "; ".join(product.testing) if product.testing else "-")
    row += 1

    row = _section(sheet, row, "Missing fields", 4)
    if result.parsed.missing_fields:
        cell = sheet.cell(row=row, column=1, value="Not present in the source document:")
        cell.font = Font(bold=True, color=RED)
        row += 1
        for name in result.parsed.missing_fields:
            sheet.cell(row=row, column=1, value=f"   - {name}").fill = PatternFill("solid", fgColor=RED_FILL)
            row += 1
    else:
        row = _kv(sheet, row, "None", "Every expected field was present.")

    if result.product.derived:
        row += 1
        row = _section(sheet, row, "Worked out automatically", 4)
        for name, basis in result.product.derived.items():
            sheet.cell(row=row, column=1, value=f"   - {name.replace('_', ' ')}")
            sheet.cell(row=row, column=2, value=basis).font = SMALL
            row += 1

    if result.parsed.parser_notes:
        row += 1
        row = _section(sheet, row, "Parser notes", 4)
        for note in result.parsed.parser_notes:
            sheet.cell(row=row, column=1, value=f"   - {note}").font = SMALL
            row += 1

    row += 1
    sheet.cell(row=row, column=1,
               value=f"Source: {result.parsed.source_name} ({result.parsed.source_format}) | "
                     f"Quote {result.quote_id} | Prepared {result.created_at:%Y-%m-%d %H:%M}").font = SMALL


# ------------------------------------------------------------- Tab 2

BOM_HEADERS = [
    "Ingredient", "Match Status", "Match Reason", "Input Part Code", "Matched Code",
    "PO Description", "Match Method", "Claimed mg/Serving", "Potency", "Overage %",
    "Overage Class", "Formula mg/Serving", "kg/Bottle", "Latest Cost $/kg",
    "Latest PO Date", "Vendor", "Vendor Count", "PO Count", "Min Cost Ever",
    "Max Cost Ever", "Cost/Bottle (Primary)", "Cost/Bottle (Low -10%)",
    "Cost/Bottle (High +10%)", "Confidence", "Notes",
]


def _bom(book: Workbook, result: QuoteResult) -> None:
    sheet = book.create_sheet("BOM - Raw Materials")
    _title(sheet, "BOM / RAW MATERIALS", len(BOM_HEADERS))
    _header_row(sheet, 3, BOM_HEADERS)
    _widths(sheet, [32, 14, 52, 14, 13, 40, 34, 15, 10, 10, 18, 16, 13, 14,
                    14, 20, 12, 10, 13, 13, 18, 18, 18, 12, 50])

    row = 4
    for line in result.ingredients:
        match = line.match
        values = [
            line.name, match.status, match.reason, _dash(line.input_part_code),
            _dash(match.matched_code), _dash(match.po_description), match.method,
            line.claimed_mg, line.potency,
            line.overage_pct, _dash(line.overage_class),
            line.formula_mg_per_serving, line.kg_per_bottle,
            _money_or_excluded(line, line.cost_per_kg),
            match.latest_po_date.isoformat() if match.latest_po_date else "-",
            _dash(match.vendor), _dash(match.unique_vendor_count), _dash(match.po_count),
            _money_or_excluded(line, match.min_unit_cost_ever),
            _money_or_excluded(line, match.max_unit_cost_ever),
            _money_or_excluded(line, line.cost_per_bottle),
            _money_or_excluded(line, line.cost_low),
            _money_or_excluded(line, line.cost_high),
            line.confidence, "; ".join(line.notes),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=column in (3, 6, 7, 25), vertical="top")
        sheet.cell(row=row, column=2).fill = STATUS_FILL.get(
            match.status, PatternFill("solid", fgColor=RED_FILL)
        )
        sheet.cell(row=row, column=24).fill = CONFIDENCE_FILL.get(line.confidence, PatternFill())
        sheet.cell(row=row, column=10).number_format = PCT
        for column in (14, 19, 20, 21, 22, 23):
            if isinstance(sheet.cell(row=row, column=column).value, (int, float)):
                sheet.cell(row=row, column=column).number_format = MONEY
        for column in (8, 12):
            sheet.cell(row=row, column=column).number_format = "#,##0.000"
        sheet.cell(row=row, column=13).number_format = "0.00000000"
        row += 1

    accepted = [line for line in result.ingredients if line.match.accepted]
    sheet.cell(row=row, column=1, value=f"TOTAL - Accepted only ({len(accepted)} of {len(result.ingredients)} lines)").font = BOLD
    for column, values in (
        (21, [line.cost_per_bottle for line in accepted]),
        (22, [line.cost_low for line in accepted]),
        (23, [line.cost_high for line in accepted]),
    ):
        cell = sheet.cell(row=row, column=column, value=sum(v or 0.0 for v in values))
        cell.font = BOLD
        cell.number_format = MONEY
        cell.fill = PatternFill("solid", fgColor=LIGHT)
    sheet.auto_filter.ref = f"A3:{get_column_letter(len(BOM_HEADERS))}{row - 1}"


# ------------------------------------------------------------- Tab 3

PKG_HEADERS = [
    "Role", "Description", "Match Status", "Match Reason", "Input Part Code",
    "Matched Code", "PO Description", "Match Method", "Qty/Bottle",
    "Bottles per Purchased Unit", "Latest Unit Cost", "Latest PO Date", "Vendor", "Min Cost Ever", "Max Cost Ever",
    "Cost/Bottle (Primary)", "Cost/Bottle (Low -10%)", "Cost/Bottle (High +10%)",
    "Confidence", "Notes",
]


def _packaging(book: Workbook, result: QuoteResult) -> None:
    sheet = book.create_sheet("Packaging")
    _title(sheet, "PACKAGING", len(PKG_HEADERS))
    _header_row(sheet, 3, PKG_HEADERS)
    _widths(sheet, [16, 36, 14, 52, 14, 13, 40, 34, 12, 14, 16, 14, 20, 13, 13,
                    18, 18, 18, 12, 50])

    row = 4
    for line in result.packaging:
        match = line.match
        values = [
            line.role.replace("_", " ").title(), line.description, match.status,
            match.reason, _dash(line.input_part_code), _dash(match.matched_code),
            _dash(match.po_description), match.method, line.qty_per_bottle,
            _dash(line.units_per_container),
            _money_or_excluded(line, line.unit_cost),
            match.latest_po_date.isoformat() if match.latest_po_date else "-",
            _dash(match.vendor),
            _money_or_excluded(line, match.min_unit_cost_ever),
            _money_or_excluded(line, match.max_unit_cost_ever),
            _money_or_excluded(line, line.cost_per_bottle),
            _money_or_excluded(line, line.cost_low),
            _money_or_excluded(line, line.cost_high),
            line.confidence, "; ".join(line.notes),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=column in (4, 7, 8, 20), vertical="top")
        sheet.cell(row=row, column=3).fill = STATUS_FILL.get(
            match.status, PatternFill("solid", fgColor=RED_FILL)
        )
        sheet.cell(row=row, column=19).fill = CONFIDENCE_FILL.get(line.confidence, PatternFill())
        for column in (11, 14, 15, 16, 17, 18):
            if isinstance(sheet.cell(row=row, column=column).value, (int, float)):
                sheet.cell(row=row, column=column).number_format = MONEY
        row += 1

    accepted = [line for line in result.packaging if line.match.accepted]
    sheet.cell(row=row, column=1, value=f"TOTAL - Accepted only ({len(accepted)} of {len(result.packaging)} lines)").font = BOLD
    for column, values in (
        (16, [line.cost_per_bottle for line in accepted]),
        (17, [line.cost_low for line in accepted]),
        (18, [line.cost_high for line in accepted]),
    ):
        cell = sheet.cell(row=row, column=column, value=sum(v or 0.0 for v in values))
        cell.font = BOLD
        cell.number_format = MONEY
        cell.fill = PatternFill("solid", fgColor=LIGHT)
    sheet.auto_filter.ref = f"A3:{get_column_letter(len(PKG_HEADERS))}{row - 1}"


# ------------------------------------------------------------- Tab 4

def _cost_summary(book: Workbook, result: QuoteResult) -> None:
    sheet = book.create_sheet("Cost Summary")
    _widths(sheet, [38, 20, 16, 16, 14])
    _title(sheet, "COST SUMMARY", 5)
    summary = result.summary
    manufacturing = result.manufacturing

    row = 3
    row = _section(sheet, row, "Estimated cost per bottle (Accepted lines only)", 5)
    row = _kv(sheet, row, "Primary", summary.primary_per_bottle, MONEY2)
    row = _kv(sheet, row, "Worst case low (-10% materials)", summary.low_per_bottle, MONEY2)
    row = _kv(sheet, row, "Worst case high (+10% materials)", summary.high_per_bottle, MONEY2)
    if summary.excluded_note:
        cell = sheet.cell(row=row, column=1, value=summary.excluded_note)
        cell.font = Font(bold=True, color=RED)
        cell.fill = PatternFill("solid", fgColor=RED_FILL)
        row += 1
    confidence_row = row
    row = _kv(sheet, row, "Quote confidence", summary.quote_confidence)
    sheet.cell(row=confidence_row, column=2).fill = CONFIDENCE_FILL.get(
        summary.quote_confidence, PatternFill()
    )
    sheet.cell(row=confidence_row, column=2).font = BOLD
    row += 1

    row = _section(sheet, row, "Cost breakdown", 5)
    for label, value in (
        ("Raw materials", summary.raw_materials),
        ("Packaging", summary.packaging),
        ("Manufacturing", summary.manufacturing),
        ("Testing", summary.testing),
    ):
        row = _kv(sheet, row, label, value, MONEY2)
    sheet.cell(row=row, column=1, value="Total").font = BOLD
    total = sheet.cell(row=row, column=2, value=summary.primary_per_bottle)
    total.font = BOLD
    total.number_format = MONEY2
    row += 2

    row = _section(sheet, row, "Cost exposure - excluded from the primary total", 5)
    exposure = _kv(sheet, row, "Needs Review lines", summary.needs_review_count)
    row = exposure
    row = _kv(sheet, row, "Tentative cost from Needs Review", summary.tentative_exposure, MONEY)
    row = _kv(sheet, row, "Unmatched lines", summary.unmatched_count)
    row = _kv(sheet, row, "Unmatched items",
              "; ".join(summary.unmatched_items) if summary.unmatched_items else "-")
    row = _kv(sheet, row, "Needs Review items",
              "; ".join(summary.needs_review_items) if summary.needs_review_items else "-")
    if summary.unmatched_count or summary.needs_review_count:
        warning = sheet.cell(
            row=row, column=1,
            value="The primary total understates true cost until these lines are resolved.",
        )
        warning.font = Font(bold=True, color=RED)
        warning.fill = PatternFill("solid", fgColor=RED_FILL)
        row += 1
    row += 1

    row = _section(sheet, row, "Top 10 cost drivers", 5)
    _header_row(sheet, row, ["#", "Driver", "Cost/Bottle", "Share of Primary", ""])
    row += 1
    for index, driver in enumerate(summary.cost_drivers, start=1):
        sheet.cell(row=row, column=1, value=index)
        sheet.cell(row=row, column=2, value=driver.label)
        cost = sheet.cell(row=row, column=3, value=driver.cost_per_bottle)
        cost.number_format = MONEY
        share = sheet.cell(row=row, column=4, value=driver.share_pct / 100.0)
        share.number_format = PCT
        row += 1
    row += 1

    if result.price_breaks:
        row = _section(sheet, row, "Volume price breaks", 5)
        _header_row(sheet, row, ["Bottles", "Cost/Bottle", "Raw Materials",
                                 "Packaging", "Manufacturing"])
        row += 1
        for item in result.price_breaks:
            sheet.cell(row=row, column=1, value=item.volume_bottles).number_format = "#,##0"
            for column, value in ((2, item.primary_per_bottle), (3, item.raw_materials),
                                  (4, item.packaging), (5, item.manufacturing)):
                cell = sheet.cell(row=row, column=column, value=value)
                cell.number_format = MONEY2
            if item.is_quoted_volume:
                for column in range(1, 6):
                    sheet.cell(row=row, column=column).fill = PatternFill("solid", fgColor=LIGHT)
                    sheet.cell(row=row, column=column).font = BOLD
            row += 1
        sheet.cell(row=row, column=1,
                   value="Material cost per bottle is flat across the ladder: PO history "
                         "carries no volume-tiered pricing.").font = SMALL
        row += 2

    if result.pricing:
        row = _section(sheet, row, "Recommended price - requires Sales and Finance review", 5)
        _header_row(sheet, row, ["Channel", "Target Margin", "Price/Bottle",
                                 "Margin/Bottle", "Basis"])
        row += 1
        for item in result.pricing:
            sheet.cell(row=row, column=1, value=item.label)
            sheet.cell(row=row, column=2, value=item.target_margin_pct / 100.0).number_format = PCT
            sheet.cell(row=row, column=3, value=item.price_per_bottle).number_format = MONEY2
            sheet.cell(row=row, column=4, value=item.margin_dollars).number_format = MONEY2
            cell = sheet.cell(row=row, column=5, value=item.basis)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        warning = sheet.cell(
            row=row, column=1,
            value="Margin targets are configured reference data, not a Finance decision. "
                  "Quick Quote does not set final pricing, margin or customer-facing terms.")
        warning.font = Font(bold=True, color=RED)
        warning.fill = PatternFill("solid", fgColor=AMBER_FILL)
        row += 2

    row = _section(sheet, row, f"Manufacturing route - {manufacturing.dosage_form or 'capsule'}", 5)
    if manufacturing.steps:
        _header_row(sheet, row, ["Step", "Work Centre", "Hours", "$/hr", "$/Bottle"])
        row += 1
        for step in manufacturing.steps:
            sheet.cell(row=row, column=1, value=step.step)
            sheet.cell(row=row, column=2, value=step.work_centre or "-")
            sheet.cell(row=row, column=3, value=step.hours).number_format = "0.00"
            if step.rate_per_hour is not None:
                sheet.cell(row=row, column=4, value=step.rate_per_hour).number_format = MONEY2
            if step.cost_per_bottle is None:
                cell = sheet.cell(row=row, column=5, value="Not costed")
                cell.fill = PatternFill("solid", fgColor=RED_FILL)
                note = sheet.cell(row=row, column=6, value=step.reason)
                note.alignment = Alignment(wrap_text=True, vertical="top")
                note.font = SMALL
            else:
                sheet.cell(row=row, column=5, value=step.cost_per_bottle).number_format = MONEY
            row += 1
        row += 1

    if manufacturing.testing_basis:
        row = _kv(sheet, row, "Testing $/bottle", manufacturing.testing_per_bottle, MONEY)
        row = _kv(sheet, row, "Testing basis", manufacturing.testing_basis)
        row += 1

    if manufacturing.estimated:
        row = _kv(sheet, row, "Components in formula", manufacturing.component_count)
        row = _kv(sheet, row, "Bottles in run", manufacturing.bottles_in_run)
        row = _kv(sheet, row, "Compounding hours", manufacturing.compounding_hours)
        row = _kv(sheet, row, "Compounding $/bottle", manufacturing.compounding_per_bottle, MONEY)
        row = _kv(sheet, row, "Encapsulation $/bottle", manufacturing.encapsulation_per_bottle, MONEY)
        row = _kv(sheet, row, "Bottling $/bottle", manufacturing.bottling_per_bottle, MONEY)
        row = _kv(sheet, row, "Capsules in run", manufacturing.total_capsules)
        row = _kv(sheet, row, "Machine assumption",
                  f"{manufacturing.machine}"
                  + (f" + {manufacturing.bottling_line}" if manufacturing.bottling_line else "")
                  + (" (selected)" if manufacturing.machine_assumed else ""))
        row = _kv(sheet, row, "Machine selection basis", manufacturing.machine_basis)
        row = _kv(sheet, row, "Basis", manufacturing.reason)
    else:
        cell = sheet.cell(row=row, column=1, value=f"Not estimated - {manufacturing.reason}")
        cell.font = Font(bold=True, color=RED)
        cell.fill = PatternFill("solid", fgColor=RED_FILL)


# ------------------------------------------------------------- Tab 5

def _review_flags(book: Workbook, result: QuoteResult) -> None:
    sheet = book.create_sheet("Review Flags")
    _widths(sheet, [16, 40, 96, 12])
    _title(sheet, "REVIEW FLAGS", 4)
    _header_row(sheet, 3, ["Owner", "Item", "Reason", "Severity"])

    row = 4
    grouped = result.flags_by_owner()
    for owner in FLAG_OWNERS:
        flags = grouped.get(owner, [])
        header = sheet.cell(row=row, column=1, value=owner)
        header.font = BOLD
        header.fill = PatternFill("solid", fgColor=LIGHT)
        count = sheet.cell(row=row, column=2, value=f"{len(flags)} flag(s)")
        count.fill = PatternFill("solid", fgColor=LIGHT)
        for column in (3, 4):
            sheet.cell(row=row, column=column).fill = PatternFill("solid", fgColor=LIGHT)
        row += 1

        if not flags:
            sheet.cell(row=row, column=2, value="No flags for this owner.").font = SMALL
            row += 1
        for flag in flags:
            sheet.cell(row=row, column=2, value=flag.item).alignment = Alignment(vertical="top", wrap_text=True)
            reason = sheet.cell(row=row, column=3, value=flag.reason)
            reason.alignment = Alignment(wrap_text=True, vertical="top")
            severity = sheet.cell(row=row, column=4, value=flag.severity)
            if flag.severity == "blocking":
                severity.font = Font(bold=True, color=RED)
                severity.fill = PatternFill("solid", fgColor=RED_FILL)
            for column in range(1, 5):
                sheet.cell(row=row, column=column).border = BORDER
            row += 1
        row += 1


# ------------------------------------------------------------- Tab 6

def _customer_summary(book: Workbook, result: QuoteResult) -> None:
    sheet = book.create_sheet("Customer Quote Summary")
    _widths(sheet, [34, 30, 18, 16])
    _title(sheet, "CUSTOMER QUOTE SUMMARY", 4)

    draft = sheet.cell(row=2, column=1, value="INTERNAL DRAFT - not for release until Sales and Finance review")
    draft.font = Font(bold=True, color="FFFFFF")
    draft.fill = PatternFill("solid", fgColor=RED)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)

    product = result.product
    summary = result.summary
    row = 4
    row = _kv(sheet, row, "Customer", _dash(product.customer))
    row = _kv(sheet, row, "Product", _dash(product.formula_name))
    row = _kv(sheet, row, "Form", " ".join(
        part for part in [
            product.dosage_form,
            f"size {product.capsule_size}" if product.capsule_size else None,
            product.capsule_type,
            f"{product.count_per_bottle} ct" if product.count_per_bottle else None,
            f"{product.servings_per_bottle} servings" if product.servings_per_bottle else None,
        ] if part) or "-")
    row = _kv(sheet, row, "Volume (bottles)", _dash(product.annual_volume_bottles))
    row += 1

    row = _section(sheet, row, "Estimated cost per bottle (Accepted lines only)", 4)
    row = _kv(sheet, row, "Primary", summary.primary_per_bottle, MONEY2)
    row = _kv(sheet, row, "Range", f"${summary.low_per_bottle:,.2f} - ${summary.high_per_bottle:,.2f}")
    confidence_row = row
    row = _kv(sheet, row, "Confidence", summary.quote_confidence)
    sheet.cell(row=confidence_row, column=2).fill = CONFIDENCE_FILL.get(
        summary.quote_confidence, PatternFill()
    )
    sheet.cell(row=confidence_row, column=2).font = BOLD
    row += 1

    if summary.unmatched_count or summary.needs_review_count:
        banner = sheet.cell(row=row, column=1, value="COST EXPOSURE")
        banner.font = Font(bold=True, color="FFFFFF")
        banner.fill = PatternFill("solid", fgColor=RED)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1
        for text in (
            f"{summary.unmatched_count} unmatched line(s) have no cost: "
            f"{'; '.join(summary.unmatched_items) or 'none'}",
            f"{summary.needs_review_count} line(s) await Purchasing confirmation, carrying "
            f"${summary.tentative_exposure:,.4f}/bottle excluded from the total.",
            "The estimate above understates true cost until these are resolved.",
        ):
            cell = sheet.cell(row=row, column=1, value=text)
            cell.fill = PatternFill("solid", fgColor=RED_FILL)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            sheet.row_dimensions[row].height = 28
            row += 1
        row += 1

    row = _section(sheet, row, "Cost breakdown", 4)
    for label, value in (
        ("Raw materials", summary.raw_materials),
        ("Packaging", summary.packaging),
        ("Manufacturing", summary.manufacturing),
        ("Testing", summary.testing),
    ):
        row = _kv(sheet, row, label, value, MONEY2)
    row += 1

    row = _section(sheet, row, "Top 5 cost drivers", 4)
    for index, driver in enumerate(summary.cost_drivers[:5], start=1):
        sheet.cell(row=row, column=1, value=f"{index}. {driver.label}")
        cost = sheet.cell(row=row, column=2, value=driver.cost_per_bottle)
        cost.number_format = MONEY
        share = sheet.cell(row=row, column=3, value=driver.share_pct / 100.0)
        share.number_format = PCT
        row += 1
    row += 1

    row = _section(sheet, row, "Items requiring review", 4)
    blocking = [flag for flag in result.flags if flag.severity == "blocking"]
    if not blocking:
        sheet.cell(row=row, column=1, value="No blocking items.").font = SMALL
        row += 1
    for flag in blocking:
        sheet.cell(row=row, column=1, value=flag.owner).font = BOLD
        sheet.cell(row=row, column=2, value=flag.item)
        cell = sheet.cell(row=row, column=3, value=flag.reason)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    row += 1
    sheet.cell(row=row, column=1,
               value=f"Quote {result.quote_id} | prepared {result.created_at:%Y-%m-%d %H:%M} | "
                     "Quick Quote prepares packages for internal review; it does not set final "
                     "pricing, margin or customer-facing terms.").font = SMALL


def build_workbook(result: QuoteResult) -> bytes:
    """Render the complete six-tab workbook to bytes."""
    book = Workbook()
    book.remove(book.active)

    _product_summary(book, result)
    _bom(book, result)
    _packaging(book, result)
    _cost_summary(book, result)
    _review_flags(book, result)
    _customer_summary(book, result)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
