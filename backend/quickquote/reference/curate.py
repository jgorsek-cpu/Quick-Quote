"""Curation worksheet: take the reference tables to R&D and bring them back.

Nine hundred proposed identities is not nine hundred decisions. A handful of
parts carry most of the purchasing spend, so the worksheet is ranked by spend
and marked with the cumulative share, letting R&D and Purchasing stop when
the remainder stops mattering.

Where a description states its own standardisation -- ``TURMERIC EXT. 95%`` --
that figure is offered in a separate *suggested* column with the text it came
from. It is never written into the potency column and never reaches the
engine unconfirmed: a suggestion a person accepts is reference data, a
suggestion applied silently is invented data.

    python -m quickquote.reference.curate export --data-dir var/reference \\
        --out var/curation.xlsx
    python -m quickquote.reference.curate apply  --data-dir var/reference \\
        --worksheet var/curation.xlsx
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from pathlib import Path

from ..engine.identity import resolve_identity
from .loader import PoRow, ReferenceData, is_component_part, load_reference_data

# A percentage in a description usually states the standardisation the claim
# is made on: "95% curcuminoids", "80% silymarin".
_PERCENT_RE = re.compile(r"(\d{1,3}(?:\.\d+)?)\s*%")

INGREDIENT_SHEET = "Ingredients"
PACKAGING_SHEET = "Packaging"
MACHINE_SHEET = "Machines"
PRICING_SHEET = "Margins"

# Columns a human fills in. Everything else is context.
INGREDIENT_COLUMNS = [
    "Part", "Description", "PO Spend", "Cumulative %", "POs", "Latest $/unit",
    "Identity", "Status", "Overage Class *", "Potency *", "Bulk Density g/mL *",
    "Suggested Potency", "Suggestion Basis", "Notes",
]
PACKAGING_COLUMNS = [
    "Part", "Description", "PO Spend", "Cumulative %", "Identity", "Role *",
    "Size", "Bottles per Purchased Unit *", "Notes",
]


@dataclass
class CurationRow:
    row: PoRow
    identity: str | None
    overage_class: str
    potency: float | None
    density: float | None
    density_defaulted: bool
    suggested_potency: float | None
    suggestion_basis: str
    share: float
    cumulative: float


def suggest_potency(description: str) -> tuple[float | None, str]:
    """Read a standardisation percentage out of a description, if it states one.

    Returned as a suggestion with its source text, never as a fact.
    """
    match = _PERCENT_RE.search(description or "")
    if not match:
        return None, ""
    value = float(match.group(1))
    if not 0 < value <= 100:
        return None, ""
    return value / 100.0, f"description reads '{match.group(0)}'"


def _rank(rows: list[PoRow]) -> list[tuple[PoRow, float, float]]:
    """Order by purchasing spend, carrying each row's share and running total."""
    total = sum(row.total_spend for row in rows) or 1.0
    ordered = sorted(rows, key=lambda row: row.total_spend, reverse=True)
    ranked = []
    running = 0.0
    for row in ordered:
        running += row.total_spend
        ranked.append((row, row.total_spend / total, running / total))
    return ranked


def build_rows(reference: ReferenceData) -> tuple[list[CurationRow], list[CurationRow]]:
    """Split PO history into ingredient and packaging curation queues."""
    ingredients, packaging = [], []

    for rows, packaging_mode, sink in (
        ([r for r in reference.po_rows if not is_component_part(r)], False, ingredients),
        ([r for r in reference.po_rows if is_component_part(r)], True, packaging),
    ):
        for row, share, cumulative in _rank(rows):
            resolution = resolve_identity(row.description, reference, packaging=packaging_mode)
            entry = resolution.entry if resolution else None
            density, defaulted = (None, True)
            if not packaging_mode:
                density, defaulted = reference.density_for(
                    resolution.canonical if resolution else None,
                    entry.overage_class if entry else None,
                )
            suggested, basis = suggest_potency(row.description)
            sink.append(
                CurationRow(
                    row=row,
                    identity=resolution.canonical if resolution else None,
                    overage_class=entry.overage_class if entry else "",
                    potency=entry.default_potency if entry else None,
                    density=density,
                    density_defaulted=defaulted,
                    suggested_potency=suggested,
                    suggestion_basis=basis,
                    share=share,
                    cumulative=cumulative,
                )
            )
    return ingredients, packaging


# ------------------------------------------------------------- export

def export(reference: ReferenceData, destination: Path, limit: int | None = None) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    ingredients, packaging = build_rows(reference)
    if limit:
        ingredients = ingredients[:limit]
        packaging = packaging[:limit]

    book = Workbook()
    book.remove(book.active)
    navy, light, amber = "1F3864", "D9E2F3", "FFF2CC"
    head_font = Font(bold=True, color="FFFFFF")

    def header(sheet, columns):
        for index, name in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=index, value=name)
            cell.font = head_font
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 30
        sheet.freeze_panes = "A2"

    # -- ingredients ---------------------------------------------------
    sheet = book.create_sheet(INGREDIENT_SHEET)
    header(sheet, INGREDIENT_COLUMNS)
    for width, letter in zip([14, 46, 14, 12, 7, 13, 30, 12, 20, 11, 18, 16, 30, 28],
                             (get_column_letter(i) for i in range(1, 15))):
        sheet.column_dimensions[letter].width = width

    classes = ",".join(sorted(reference.overage))
    class_rule = DataValidation(type="list", formula1=f'"{classes[:250]}"', allow_blank=True)
    sheet.add_data_validation(class_rule)

    for index, item in enumerate(ingredients, start=2):
        values = [
            item.row.part_number, item.row.description, item.row.total_spend,
            item.cumulative, item.row.po_count, item.row.latest_unit_cost,
            item.identity or "", "PROPOSED" if item.identity is None else "resolved",
            item.overage_class, item.potency,
            None if item.density_defaulted else item.density,
            item.suggested_potency, item.suggestion_basis, "",
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=index, column=column, value=value)
        sheet.cell(row=index, column=3).number_format = '"$"#,##0'
        sheet.cell(row=index, column=4).number_format = "0.0%"
        sheet.cell(row=index, column=6).number_format = '"$"#,##0.0000'
        for column in (9, 10, 11):
            sheet.cell(row=index, column=column).fill = PatternFill("solid", fgColor=amber)
        class_rule.add(sheet.cell(row=index, column=9))
        if item.cumulative <= 0.80:
            sheet.cell(row=index, column=1).fill = PatternFill("solid", fgColor=light)

    # -- packaging ------------------------------------------------------
    sheet = book.create_sheet(PACKAGING_SHEET)
    header(sheet, PACKAGING_COLUMNS)
    for width, letter in zip([14, 46, 14, 12, 30, 16, 20, 24, 28],
                             (get_column_letter(i) for i in range(1, 10))):
        sheet.column_dimensions[letter].width = width

    roles = "bottle,cap,label,neckband,cotton,desiccant,seal,shipper,capsule_shell,scoop"
    role_rule = DataValidation(type="list", formula1=f'"{roles}"', allow_blank=True)
    sheet.add_data_validation(role_rule)

    from ..engine.text import extract_size

    for index, item in enumerate(packaging, start=2):
        entry = reference.identity_entry(item.identity, packaging=True) if item.identity else None
        values = [
            item.row.part_number, item.row.description, item.row.total_spend,
            item.cumulative, item.identity or "",
            entry.role if entry else "",
            extract_size(item.row.description).describe(),
            entry.units_per_container if entry else None, "",
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=index, column=column, value=value)
        sheet.cell(row=index, column=3).number_format = '"$"#,##0'
        sheet.cell(row=index, column=4).number_format = "0.0%"
        for column in (6, 8):
            sheet.cell(row=index, column=column).fill = PatternFill("solid", fgColor=amber)
        role_rule.add(sheet.cell(row=index, column=6))
        if item.cumulative <= 0.80:
            sheet.cell(row=index, column=1).fill = PatternFill("solid", fgColor=light)

    # -- machines -------------------------------------------------------
    sheet = book.create_sheet(MACHINE_SHEET)
    header(sheet, ["Machine", "Work Centre", "Labor $/hr", "Overhead $/hr",
                   "Capsules per Hour *", "Min Capsules", "Max Capsules", "Notes"])
    for width, letter in zip([18, 14, 13, 15, 20, 15, 15, 40], "ABCDEFGH"):
        sheet.column_dimensions[letter].width = width
    for index, machine in enumerate(reference.machines, start=2):
        for column, value in enumerate([
            machine.machine, machine.work_centre, machine.labor_rate_per_hour,
            machine.overhead_rate_per_hour, machine.capsules_per_hour,
            machine.min_capsules, machine.max_capsules, machine.notes,
        ], start=1):
            sheet.cell(row=index, column=column, value=value)
        sheet.cell(row=index, column=5).fill = PatternFill("solid", fgColor=amber)
    sheet.cell(row=len(reference.machines) + 3, column=1,
               value="Capsules per hour are placeholders scaled from one calibrated "
                     "rate. Replace them from the CVC run rate sheet.").font = Font(italic=True)

    # -- margins --------------------------------------------------------
    sheet = book.create_sheet(PRICING_SHEET)
    header(sheet, ["Channel", "Label", "Target Margin % *", "Notes"])
    for width, letter in zip([16, 30, 20, 60], "ABCD"):
        sheet.column_dimensions[letter].width = width
    for index, target in enumerate(reference.pricing, start=2):
        for column, value in enumerate(
            [target.channel, target.label, target.target_margin_pct, target.notes], start=1
        ):
            sheet.cell(row=index, column=column, value=value)
        sheet.cell(row=index, column=3).fill = PatternFill("solid", fgColor=amber)
    sheet.cell(row=len(reference.pricing) + 3, column=1,
               value="Margin = (price - cost) / price. Finance owns these numbers.").font = Font(italic=True)

    destination.parent.mkdir(parents=True, exist_ok=True)
    book.save(destination)

    covered = sum(1 for item in ingredients if item.cumulative <= 0.80)
    return {
        "ingredients": len(ingredients),
        "packaging": len(packaging),
        "ingredients_to_80pct": covered,
    }


# -------------------------------------------------------------- apply

def _changed(new_value, current_value) -> bool:
    """True when a worksheet cell actually answers something.

    A blank never changes anything. A value identical to what the tables
    already hold is not an answer either: the export pre-fills the current
    value as context, so treating every filled cell as a confirmation would
    silently stamp hundreds of class-level defaults as "confirmed by R&D".
    """
    if new_value in (None, ""):
        return False
    try:
        return abs(float(new_value) - float(current_value)) > 1e-12
    except (TypeError, ValueError):
        return str(new_value).strip().lower() != str(current_value or "").strip().lower()


def apply(worksheet: Path, data_dir: Path) -> dict:
    """Write confirmed answers back into the reference tables.

    Only cells that differ from the exported value are written, so re-applying
    an untouched worksheet changes nothing.
    """
    from openpyxl import load_workbook

    book = load_workbook(worksheet, data_only=True)
    changed = {"potency": 0, "overage_class": 0, "density": 0,
               "role": 0, "units_per_container": 0, "machines": 0, "margins": 0}

    reference = load_reference_data(data_dir)

    identity_path = data_dir / "ingredient_identity.csv"
    identities = {row["canonical"]: row for row in csv.DictReader(identity_path.open())}
    density_path = data_dir / "bulk_density.csv"
    densities = {row["key"]: row for row in csv.DictReader(density_path.open())}

    def cells(sheet_name):
        sheet = book[sheet_name]
        index = {name: i for i, name in enumerate([cell.value for cell in sheet[1]]) if name}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            def get(name):
                position = index.get(name)
                return row[position] if position is not None and position < len(row) else None
            yield get

    if INGREDIENT_SHEET in book.sheetnames:
        for get in cells(INGREDIENT_SHEET):
            identity = (get("Identity") or "").strip().lower()
            if not identity or identity not in identities:
                continue
            record = identities[identity]

            potency = get("Potency *")
            if _changed(potency, record.get("default_potency")):
                record["default_potency"] = f"{float(potency):g}"
                changed["potency"] += 1

            overage = get("Overage Class *")
            if _changed(overage, record.get("overage_class")):
                record["overage_class"] = str(overage).strip().lower()
                changed["overage_class"] += 1

            density = get("Bulk Density g/mL *")
            # Compare against whatever the engine would use today, so accepting
            # an inherited class default is correctly read as "no change".
            current, _ = reference.density_for(identity, record.get("overage_class"))
            if _changed(density, current):
                densities[identity] = {
                    "key": identity,
                    "bulk_density_g_ml": f"{float(density):g}",
                    "notes": "Confirmed by R&D",
                }
                changed["density"] += 1

        _write(identity_path,
               ["canonical", "aliases", "overage_class", "default_potency", "notes"],
               identities.values())
        _write(density_path, ["key", "bulk_density_g_ml", "notes"], densities.values())

    packaging_path = data_dir / "packaging_identity.csv"
    packaging = {row["canonical"]: row for row in csv.DictReader(packaging_path.open())}
    if PACKAGING_SHEET in book.sheetnames:
        for get in cells(PACKAGING_SHEET):
            identity = (get("Identity") or "").strip().lower()
            if not identity or identity not in packaging:
                continue
            record = packaging[identity]

            role = get("Role *")
            if _changed(role, record.get("role")):
                record["role"] = str(role).strip().lower()
                changed["role"] += 1

            per_container = get("Bottles per Purchased Unit *")
            if _changed(per_container, record.get("units_per_container")):
                record["units_per_container"] = f"{float(per_container):g}"
                changed["units_per_container"] += 1

        _write(packaging_path,
               ["canonical", "aliases", "role", "size_required", "units_per_container", "notes"],
               packaging.values())

    if MACHINE_SHEET in book.sheetnames:
        path = data_dir / "machines.csv"
        machines = {row["machine"]: row for row in csv.DictReader(path.open())}
        for row in book[MACHINE_SHEET].iter_rows(min_row=2, values_only=True):
            name = (row[0] or "").strip()
            if name in machines and _changed(row[4], machines[name].get("capsules_per_hour")):
                machines[name]["capsules_per_hour"] = f"{float(row[4]):g}"
                machines[name]["notes"] = "Confirmed run rate"
                changed["machines"] += 1
        _write(path, ["machine", "work_centre", "labor_rate_per_hour",
                      "overhead_rate_per_hour", "capsules_per_hour",
                      "min_capsules", "max_capsules", "notes"], machines.values())

    if PRICING_SHEET in book.sheetnames:
        path = data_dir / "pricing.csv"
        pricing = {row["channel"]: row for row in csv.DictReader(path.open())}
        for row in book[PRICING_SHEET].iter_rows(min_row=2, values_only=True):
            channel = (row[0] or "").strip().lower()
            if channel in pricing and _changed(row[2], pricing[channel].get("target_margin_pct")):
                pricing[channel]["target_margin_pct"] = f"{float(row[2]):g}"
                pricing[channel]["notes"] = "Confirmed by Finance"
                changed["margins"] += 1
        _write(path, ["channel", "target_margin_pct", "label", "notes"], pricing.values())

    return changed


def _write(path: Path, fields: list[str], rows) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m quickquote.reference.curate",
        description="Take the reference tables to R&D and bring the answers back.",
    )
    sub = parser.add_subparsers(dest="command", required=True)

    out = sub.add_parser("export", help="Write a curation worksheet ranked by spend.")
    out.add_argument("--data-dir", type=Path, required=True)
    out.add_argument("--out", type=Path, required=True)
    out.add_argument("--limit", type=int, default=None,
                     help="Only the top N rows of each queue.")

    back = sub.add_parser("apply", help="Write confirmed answers back.")
    back.add_argument("--data-dir", type=Path, required=True)
    back.add_argument("--worksheet", type=Path, required=True)

    arguments = parser.parse_args(argv)

    if arguments.command == "export":
        reference = load_reference_data(arguments.data_dir)
        stats = export(reference, arguments.out, arguments.limit)
        print("Quick Quote curation worksheet")
        print(f"  ingredient rows: {stats['ingredients']:,}")
        print(f"  packaging rows:  {stats['packaging']:,}")
        print(f"  ingredients covering 80% of spend: {stats['ingredients_to_80pct']:,} "
              "(shaded in the Part column)")
        print(f"\nWrote {arguments.out}")
        print("Amber cells are the ones to fill in. Blank means 'not answered' and "
              "never clears an existing value.")
        return 0

    changed = apply(arguments.worksheet, arguments.data_dir)
    print("Applied to " + str(arguments.data_dir))
    for name, count in changed.items():
        print(f"  {name.replace('_', ' ')}: {count}")
    print("\nReload a running instance with: curl -X POST <host>/api/reference/reload")
    return 0


if __name__ == "__main__":
    sys.exit(main())
